from typing import Tuple, Dict, Any
from docx import Document

class DOCXExtractor:
    async def extract(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text and metadata from a DOCX file."""
        doc = Document(file_path)
        content = "\n".join([para.text for para in doc.paragraphs])
        metadata = {'paragraph_count': len(doc.paragraphs)}
        return content, metadata
